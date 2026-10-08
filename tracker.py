#!/usr/bin/env python3
"""Internship tracker: checks organizations' job boards, updates a spreadsheet,
and sends alerts when new internship/fellowship postings appear.

Usage:
    python tracker.py              # check every organization, update data/, send alerts
    python tracker.py --no-alerts  # update the spreadsheet only
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import smtplib
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from email.message import EmailMessage
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin

import requests
import yaml

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
POSTINGS_CSV = DATA / "postings.csv"
POSTINGS_XLSX = DATA / "internships.xlsx"
NEW_POSTINGS_MD = DATA / "new_postings.md"
RUN_REPORT_MD = DATA / "run_report.md"

HEADERS = {"User-Agent": "Mozilla/5.0 (internship-tracker; personal job search)"}
TIMEOUT = 30

COLUMNS = [
    "id", "organization", "category", "title", "location", "url", "source",
    "first_seen", "last_seen", "still_listed", "my_status", "my_notes",
]
# Columns you can edit by hand in data/postings.csv; the tracker never overwrites them.
USER_COLUMNS = ("my_status", "my_notes")


@dataclass
class Posting:
    organization: str
    category: str
    title: str
    url: str
    location: str = ""
    source: str = ""
    id: str = field(default="")

    def __post_init__(self):
        self.title = " ".join(self.title.split())
        if not self.id:
            key = f"{self.organization}|{self.url}|{self.title}".lower()
            self.id = hashlib.sha1(key.encode()).hexdigest()[:12]


# --------------------------------------------------------------------------- fetchers

def get_json(url: str):
    resp = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    resp.raise_for_status()
    return resp.json()


def get_text(url: str) -> str:
    resp = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    resp.raise_for_status()
    return resp.text


def fetch_greenhouse(org: dict) -> list[Posting]:
    data = get_json(f"https://boards-api.greenhouse.io/v1/boards/{org['source']}/jobs")
    return [
        Posting(org["name"], org["category"], j["title"], j["absolute_url"],
                (j.get("location") or {}).get("name", ""), "greenhouse")
        for j in data.get("jobs", [])
    ]


def fetch_lever(org: dict) -> list[Posting]:
    data = get_json(f"https://api.lever.co/v0/postings/{org['source']}?mode=json")
    return [
        Posting(org["name"], org["category"], j["text"], j["hostedUrl"],
                (j.get("categories") or {}).get("location", ""), "lever")
        for j in data
    ]


def fetch_ashby(org: dict) -> list[Posting]:
    data = get_json(f"https://api.ashbyhq.com/posting-api/job-board/{org['source']}")
    return [
        Posting(org["name"], org["category"], j["title"], j["jobUrl"],
                j.get("location", ""), "ashby")
        for j in data.get("jobs", [])
    ]


def fetch_workable(org: dict) -> list[Posting]:
    data = get_json(f"https://apply.workable.com/api/v1/widget/accounts/{org['source']}")
    return [
        Posting(org["name"], org["category"], j["title"], j["url"],
                ", ".join(filter(None, [j.get("city"), j.get("country")])), "workable")
        for j in data.get("jobs", [])
    ]


class _PageParser(HTMLParser):
    """Collects link texts (with hrefs) and heading/list-item texts from a careers page."""

    BLOCKS = {"h1", "h2", "h3", "h4", "h5", "li", "p", "td"}
    SKIP = {"script", "style", "noscript", "nav", "footer"}

    def __init__(self):
        super().__init__()
        self.items: list[tuple[str, str | None]] = []  # (text, href or None)
        self._link: list | None = None
        self._block: list | None = None
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "a":
            self._link = [dict(attrs).get("href"), []]
        elif tag in self.BLOCKS:
            self._block = []

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag == "a" and self._link is not None:
            self.items.append((" ".join(self._link[1]), self._link[0]))
            self._link = None
        elif tag in self.BLOCKS and self._block is not None:
            self.items.append((" ".join(self._block), None))
            self._block = None

    def handle_data(self, data):
        if self._skip or not data.strip():
            return
        if self._link is not None:
            self._link[1].append(data.strip())
        if self._block is not None:
            self._block.append(data.strip())


def fetch_page(org: dict) -> list[Posting]:
    """Generic careers page: every link or heading/list item becomes a candidate posting.

    Keyword filtering happens later, so only lines mentioning e.g. "internship" survive.
    """
    url = org["source"]
    parser = _PageParser()
    parser.feed(get_text(url))
    postings, seen = [], set()
    for text, href in parser.items:
        text = " ".join(text.split())
        if not text or len(text) > 160:  # long paragraphs are descriptions, not titles
            continue
        link = urljoin(url, href) if href and not href.startswith(("#", "mailto:", "javascript:")) else url
        if text.lower() in seen:
            continue
        seen.add(text.lower())
        postings.append(Posting(org["name"], org["category"], text, link, "", "page"))
    return postings


FETCHERS = {
    "greenhouse": fetch_greenhouse,
    "lever": fetch_lever,
    "ashby": fetch_ashby,
    "workable": fetch_workable,
    "page": fetch_page,
}


# --------------------------------------------------------------------------- filtering

def _pattern(words: list[str]) -> re.Pattern | None:
    words = [w.strip() for w in words or [] if w and w.strip()]
    if not words:
        return None
    return re.compile(r"\b(" + "|".join(re.escape(w) for w in words) + r")\b", re.IGNORECASE)


def make_filter(config: dict):
    include = _pattern(config.get("include_keywords"))
    exclude = _pattern(config.get("exclude_keywords"))
    locations = [l.lower() for l in config.get("locations") or []]

    def keep(p: Posting) -> bool:
        if include and not include.search(p.title):
            return False
        if exclude and exclude.search(p.title):
            return False
        if locations and p.location and not any(l in p.location.lower() for l in locations):
            return False
        return True

    return keep


# --------------------------------------------------------------------------- spreadsheet

def load_existing() -> dict[str, dict]:
    if not POSTINGS_CSV.exists():
        return {}
    with POSTINGS_CSV.open(newline="", encoding="utf-8") as f:
        return {row["id"]: row for row in csv.DictReader(f)}


def merge(existing: dict[str, dict], found: list[Posting], checked_orgs: set[str], today: str):
    """Returns (rows, new_postings). Postings from orgs that failed this run are left untouched."""
    rows = {k: dict(v) for k, v in existing.items()}
    new = []
    found_ids = set()
    for p in found:
        found_ids.add(p.id)
        if p.id in rows:
            rows[p.id].update(title=p.title, location=p.location, url=p.url,
                              last_seen=today, still_listed="Yes")
        else:
            rows[p.id] = {
                "id": p.id, "organization": p.organization, "category": p.category,
                "title": p.title, "location": p.location, "url": p.url, "source": p.source,
                "first_seen": today, "last_seen": today, "still_listed": "Yes",
                "my_status": "", "my_notes": "",
            }
            new.append(p)
    for row_id, row in rows.items():
        if row["organization"] in checked_orgs and row_id not in found_ids:
            row["still_listed"] = "No"
    # Open postings first (newest first), then ones that have been taken down.
    listed = sorted((r for r in rows.values() if r["still_listed"] == "Yes"),
                    key=lambda r: (r["first_seen"], r["organization"]), reverse=True)
    gone = sorted((r for r in rows.values() if r["still_listed"] != "Yes"),
                  key=lambda r: r["last_seen"], reverse=True)
    return listed + gone, new


def write_csv(rows: list[dict]):
    with POSTINGS_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows: list[dict], new_ids: set[str], orgs: list[dict], errors: dict[str, str]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Postings"
    ws.append([c.replace("_", " ").title() for c in COLUMNS[1:]])
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    new_fill = PatternFill("solid", fgColor="E2EFDA")
    gone_font = Font(color="999999")
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS[1:]])
        row = ws.max_row
        url_cell = ws.cell(row=row, column=COLUMNS.index("url"))
        if r["url"]:
            url_cell.hyperlink = r["url"]
            url_cell.font = Font(color="0563C1", underline="single")
        if r["id"] in new_ids:
            for cell in ws[row]:
                cell.fill = new_fill
        if r["still_listed"] != "Yes":
            for cell in ws[row]:
                cell.font = gone_font
    widths = {"organization": 30, "category": 18, "title": 55, "location": 22, "url": 45,
              "source": 11, "first_seen": 12, "last_seen": 12, "still_listed": 11,
              "my_status": 14, "my_notes": 40}
    for i, c in enumerate(COLUMNS[1:], start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 15)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    os_ = wb.create_sheet("Organizations")
    os_.append(["Name", "Category", "Source type", "Source", "Open matches", "Last check", "Notes"])
    for cell in os_[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for o in orgs:
        count = sum(1 for r in rows if r["organization"] == o["name"] and r["still_listed"] == "Yes")
        status = f"ERROR: {errors[o['name']]}" if o["name"] in errors else "OK"
        os_.append([o["name"], o["category"], o["source_type"], o["source"], count, status,
                    o.get("notes", "")])
    for col, width in zip("ABCDEFG", (40, 18, 12, 50, 13, 40, 50)):
        os_.column_dimensions[col].width = width
    for row in os_.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(POSTINGS_XLSX)


# --------------------------------------------------------------------------- alerts

def format_alert(new: list[Posting]) -> str:
    lines = [f"**{len(new)} new internship/fellowship posting(s) found**", ""]
    by_org: dict[str, list[Posting]] = {}
    for p in new:
        by_org.setdefault(p.organization, []).append(p)
    for org in sorted(by_org):
        lines.append(f"### {org}")
        for p in by_org[org]:
            loc = f" — {p.location}" if p.location else ""
            lines.append(f"- [{p.title}]({p.url}){loc}")
        lines.append("")
    lines.append("Full list: `data/postings.csv` / `data/internships.xlsx`")
    return "\n".join(lines)


def send_email(subject: str, body: str):
    host, to = os.environ.get("SMTP_HOST"), os.environ.get("ALERT_EMAIL_TO")
    if not (host and to):
        return
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = os.environ.get("SMTP_FROM") or os.environ.get("SMTP_USER") or to
    msg["To"] = to
    msg.set_content(body)
    with smtplib.SMTP(host, int(os.environ.get("SMTP_PORT", "587")), timeout=TIMEOUT) as s:
        s.starttls()
        if os.environ.get("SMTP_USER"):
            s.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASSWORD", ""))
        s.send_message(msg)
    print(f"Email alert sent to {to}")


def send_slack(body: str):
    hook = os.environ.get("SLACK_WEBHOOK_URL")
    if not hook:
        return
    requests.post(hook, json={"text": body}, timeout=TIMEOUT).raise_for_status()
    print("Slack alert sent")


def set_github_output(**values):
    path = os.environ.get("GITHUB_OUTPUT")
    if path:
        with open(path, "a") as f:
            for k, v in values.items():
                f.write(f"{k}={v}\n")


# --------------------------------------------------------------------------- main

def load_orgs(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as f:
        orgs = [
            {k: (v or "").strip() for k, v in row.items()}
            for row in csv.DictReader(f)
            if (row.get("name") or "").strip() and not row["name"].lstrip().startswith("#")
        ]
    for o in orgs:
        o["source_type"] = o["source_type"].lower()
    return orgs


def run(config_path: Path, orgs_path: Path, alerts: bool) -> int:
    config = yaml.safe_load(config_path.read_text()) or {}
    orgs = load_orgs(orgs_path)
    keep = make_filter(config)
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    found, errors, checked = [], {}, set()
    for org in orgs:
        fetch = FETCHERS.get(org["source_type"])
        if not fetch:
            errors[org["name"]] = f"unknown source_type '{org['source_type']}'"
            continue
        try:
            matches = [p for p in fetch(org) if keep(p)]
        except Exception as e:  # one broken site shouldn't stop the run
            errors[org["name"]] = f"{type(e).__name__}: {e}"[:200]
            print(f"  ! {org['name']}: {errors[org['name']]}", file=sys.stderr)
            continue
        checked.add(org["name"])
        found.extend(matches)
        print(f"  {org['name']}: {len(matches)} matching posting(s)")

    DATA.mkdir(exist_ok=True)
    rows, new = merge(load_existing(), found, checked, today)
    write_csv(rows)
    write_xlsx(rows, {p.id for p in new}, orgs, errors)

    report = [f"# Run report — {today}", "",
              f"- Organizations checked: {len(checked)}/{len(orgs)}",
              f"- Matching postings currently listed: {sum(r['still_listed'] == 'Yes' for r in rows)}",
              f"- New since last run: {len(new)}", ""]
    if errors:
        report += ["## Sources that failed (check the URL / slug in organizations.csv)", ""]
        report += [f"- **{name}**: {err}" for name, err in sorted(errors.items())]
    RUN_REPORT_MD.write_text("\n".join(report) + "\n")

    if new:
        body = format_alert(new)
        NEW_POSTINGS_MD.write_text(body + "\n")
        print("\n" + body)
        if alerts:
            subject = f"{config.get('email_subject_prefix', '[Internship Alert]')} {len(new)} new posting(s)"
            for send in (lambda: send_email(subject, body), lambda: send_slack(body)):
                try:
                    send()
                except Exception as e:
                    print(f"  ! alert failed: {e}", file=sys.stderr)
    elif NEW_POSTINGS_MD.exists():
        NEW_POSTINGS_MD.unlink()

    print(f"\n{len(new)} new posting(s); {len(errors)} source error(s). See {RUN_REPORT_MD.relative_to(ROOT)}")
    set_github_output(new_count=len(new), error_count=len(errors))
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", type=Path, default=ROOT / "config.yaml")
    ap.add_argument("--orgs", type=Path, default=ROOT / "organizations.csv")
    ap.add_argument("--no-alerts", action="store_true", help="don't send email/Slack alerts")
    args = ap.parse_args()
    sys.exit(run(args.config, args.orgs, alerts=not args.no_alerts))


if __name__ == "__main__":
    main()
